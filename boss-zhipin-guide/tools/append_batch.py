"""
append_batch.py — 将一批候选人记录追加写入 Excel（每 10 个招呼为一批）。
用法: python append_batch.py <xlsx路径> <json文件路径>
json 结构: {"batch": 3, "job": "机械工程师(机器人方向)", "records": [ {列名: 值, ...}, ... ]}
表头自动创建; 已有文件则追加行。支持多岗位共用一张表（靠"岗位"列区分）。
"""
import json
import sys
from pathlib import Path
from openpyxl import load_workbook, Workbook

COLUMNS = ["批次", "岗位", "编号", "姓名", "年龄", "学历", "院校(标签)", "岗位相关年限",
           "相关性评分", "判定", "不匹配点/备注", "记录时间"]

def main():
    xlsx = Path(sys.argv[1])
    data = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
    batch = data.get("batch")
    job = data.get("job", "")
    records = data.get("records", [])

    if xlsx.exists():
        wb = load_workbook(xlsx)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "打招呼记录"
        ws.append(COLUMNS)

    for r in records:
        row = [batch, job] + [r.get(c, "") for c in COLUMNS[2:-1]] + [r.get("记录时间", "")]
        ws.append(row)

    wb.save(xlsx)
    print(f"appended batch {batch} job={job}: {len(records)} rows -> {xlsx}")

if __name__ == "__main__":
    main()
